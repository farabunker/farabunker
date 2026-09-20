"""Unit tests for tools/rag/readers.py -- the "how do I read *this* file
format" seam ingest.py hands every source file to.

No binary fixtures are committed: PDF and DOCX fixtures are generated at
test time in `tmp_path`.

DOCX generation uses `python-docx` (`Document().add_paragraph(...)`)
directly -- straightforward. PDF generation does NOT use `pypdf`'s
`PdfWriter` to build the page: pypdf has no simple "add a page of real
text" call (its writer API only clones/manipulates existing pages/content
streams), so this module hand-assembles a minimal-but-valid PDF byte
string instead -- a handful of indirect objects (Catalog, Pages, Page,
a Helvetica font, and a content stream with a bare `BT ... Tj ET` text
block) plus an xref table pypdf can parse. Kept short deliberately; see
`tools.rag.tests._helpers.make_pdf_bytes`.
"""
from __future__ import annotations

from unittest.mock import patch

import pandas as pd
import pytest
from pypdf._page import PageObject

from tools.rag import readers
from tools.rag.tests._helpers import make_pdf_bytes


# --- _read_pdf -------------------------------------------------------------


class TestReadPdf:
    def test_extracts_one_document_per_non_empty_page_with_page_metadata(self, tmp_path):
        pdf_path = tmp_path / "two-page.pdf"
        pdf_path.write_bytes(make_pdf_bytes(["Hello page one", "Second page text"]))

        docs = readers._read_pdf(pdf_path)

        assert [d.text for d in docs] == ["Hello page one", "Second page text"]
        assert [d.metadata["page"] for d in docs] == [1, 2]

    def test_blank_pages_are_skipped(self, tmp_path):
        """A page whose extracted text is empty (after stripping) yields
        no Document at all -- the surviving page keeps its real 1-based
        page number, not a renumbered one."""
        pdf_path = tmp_path / "with-blank.pdf"
        pdf_path.write_bytes(make_pdf_bytes(["", "Only real content"]))

        docs = readers._read_pdf(pdf_path)

        assert len(docs) == 1
        assert docs[0].text == "Only real content"
        assert docs[0].metadata["page"] == 2

    def test_all_blank_pages_returns_empty_list_and_logs_a_warning(self, tmp_path, caplog):
        pdf_path = tmp_path / "all-blank.pdf"
        pdf_path.write_bytes(make_pdf_bytes([""]))

        with caplog.at_level("WARNING"):
            docs = readers._read_pdf(pdf_path)

        assert docs == []
        assert "no extractable text found in PDF" in caplog.text


# --- _read_text --------------------------------------------------------------


class TestReadText:
    def test_reads_utf8_text_file(self, tmp_path):
        path = tmp_path / "note.txt"
        path.write_text("plain text with é accent", encoding="utf-8")

        docs = readers._read_text(path)

        assert len(docs) == 1
        assert docs[0].text == "plain text with é accent"

    def test_reads_utf8_markdown_file(self, tmp_path):
        path = tmp_path / "note.md"
        path.write_text("# Heading\n\nSome body text.", encoding="utf-8")

        docs = readers._read_text(path)

        assert docs[0].text == "# Heading\n\nSome body text."

    def test_invalid_utf8_bytes_are_replaced_not_raised(self, tmp_path):
        """`errors="replace"` -- a malformed/non-UTF-8 byte never blows up
        the read, it degrades to the Unicode replacement character."""
        path = tmp_path / "bad.txt"
        path.write_bytes(b"before \xff\xfe after")

        docs = readers._read_text(path)

        assert "before " in docs[0].text
        assert "after" in docs[0].text


# --- _read_docx --------------------------------------------------------------


class TestReadDocx:
    def test_extracts_non_empty_paragraphs_joined_by_newline(self, tmp_path):
        import docx

        d = docx.Document()
        d.add_paragraph("First paragraph.")
        d.add_paragraph("")  # blank paragraph -- dropped
        d.add_paragraph("Second paragraph.")
        path = tmp_path / "doc.docx"
        d.save(str(path))

        docs = readers._read_docx(path)

        assert len(docs) == 1
        assert docs[0].text == "First paragraph.\nSecond paragraph."


# --- read_prose_documents (public entry point, routes by extension) ---------


class TestReadProseDocuments:
    def test_routes_pdf_to_read_pdf(self, tmp_path):
        path = tmp_path / "x.pdf"
        path.write_bytes(make_pdf_bytes(["pdf content"]))

        docs = readers.read_prose_documents(path)

        assert docs[0].text == "pdf content"

    @pytest.mark.parametrize("ext", [".txt", ".md"])
    def test_routes_txt_and_md_to_read_text(self, tmp_path, ext):
        path = tmp_path / f"x{ext}"
        path.write_text("plain text body", encoding="utf-8")

        docs = readers.read_prose_documents(path)

        assert docs[0].text == "plain text body"

    def test_routes_docx_to_read_docx(self, tmp_path):
        import docx

        d = docx.Document()
        d.add_paragraph("docx body")
        path = tmp_path / "x.docx"
        d.save(str(path))

        docs = readers.read_prose_documents(path)

        assert docs[0].text == "docx body"

    def test_extension_matching_is_case_insensitive(self, tmp_path):
        path = tmp_path / "x.TXT"
        path.write_text("upper ext body", encoding="utf-8")

        docs = readers.read_prose_documents(path)

        assert docs[0].text == "upper ext body"

    def test_unsupported_extension_raises_value_error(self, tmp_path):
        path = tmp_path / "x.rtf"
        path.write_text("body", encoding="utf-8")

        with pytest.raises(ValueError, match="Unsupported prose extension: .rtf"):
            readers.read_prose_documents(path)


# --- read_tabular_dataframe (public entry point, routes by extension) ------


class TestReadTabularDataframe:
    def test_routes_csv_to_pandas_read_csv(self, tmp_path):
        path = tmp_path / "x.csv"
        path.write_text("a,b\n1,2\n", encoding="utf-8")

        df = readers.read_tabular_dataframe(path)

        assert list(df.columns) == ["a", "b"]
        assert df.iloc[0].to_dict() == {"a": 1, "b": 2}

    def test_routes_xlsx_to_pandas_read_excel(self, tmp_path):
        path = tmp_path / "x.xlsx"
        pd.DataFrame({"a": [1], "b": [2]}).to_excel(path, index=False)

        df = readers.read_tabular_dataframe(path)

        assert list(df.columns) == ["a", "b"]
        assert df.iloc[0].to_dict() == {"a": 1, "b": 2}

    def test_unsupported_extension_raises_value_error(self, tmp_path):
        path = tmp_path / "x.parquet"
        path.write_text("body", encoding="utf-8")

        with pytest.raises(ValueError, match="Unsupported tabular extension: .parquet"):
            readers.read_tabular_dataframe(path)


# --- media extension registry -----------------------------------------------


class TestMediaExtensionSets:
    def test_video_exts(self):
        assert readers.VIDEO_EXTS == {".mp4", ".mov", ".mkv", ".webm"}

    def test_audio_exts(self):
        assert readers.AUDIO_EXTS == {".mp3", ".wav", ".m4a"}

    def test_image_exts_includes_jpeg(self):
        assert readers.IMAGE_EXTS == {".png", ".jpg", ".jpeg"}

    def test_av_exts_is_video_union_audio(self):
        assert readers.AV_EXTS == readers.VIDEO_EXTS | readers.AUDIO_EXTS

    def test_media_exts_is_av_union_image(self):
        assert readers.MEDIA_EXTS == readers.AV_EXTS | readers.IMAGE_EXTS


# --- medium_for --------------------------------------------------------------


class TestMediumFor:
    @pytest.mark.parametrize("ext", sorted(readers.PROSE_EXTS))
    def test_prose_extensions_route_to_prose(self, ext):
        assert readers.medium_for(ext) == "prose"

    @pytest.mark.parametrize("ext", sorted(readers.TABULAR_EXTS))
    def test_tabular_extensions_route_to_tabular(self, ext):
        assert readers.medium_for(ext) == "tabular"

    @pytest.mark.parametrize("ext", sorted(readers.VIDEO_EXTS))
    def test_video_extensions_route_to_video(self, ext):
        assert readers.medium_for(ext) == "video"

    @pytest.mark.parametrize("ext", sorted(readers.AUDIO_EXTS))
    def test_audio_extensions_route_to_audio(self, ext):
        assert readers.medium_for(ext) == "audio"

    @pytest.mark.parametrize("ext", sorted(readers.IMAGE_EXTS))
    def test_image_extensions_route_to_image(self, ext):
        assert readers.medium_for(ext) == "image"

    def test_case_insensitive(self):
        assert readers.medium_for(".PDF") == "prose"
        assert readers.medium_for(".MP4") == "video"
        assert readers.medium_for(".JPG") == "image"

    def test_unsupported_extension_raises_value_error_with_supported_list(self):
        with pytest.raises(ValueError, match=r"Unsupported file extension: '\.rtf' \(supported: \[") as exc_info:
            readers.medium_for(".rtf")
        assert ".pdf" in str(exc_info.value)
        assert ".mp4" in str(exc_info.value)


# --- pdf_textless_pages (W1) -------------------------------------------------


class TestPdfTextlessPages:
    """`readers.pdf_textless_pages` replaces the retired `pdf_has_extractable_
    text`/`pdf_page_count` pair (W1, D2/D1) -- the per-page scanned/mixed-PDF
    detector: 1-based page numbers whose `extract_text()` is blank, with an
    optional `limit` bounding how many textless pages get COLLECTED (not
    scanned) before the scan stops early.

    RETARGETED (H28 review round 1, finding 4): this function ALWAYS
    returns `(pages, truncated)` now, never a bare list -- every assertion
    below unpacks or compares against the pair. None of these calls pass
    `max_examined`, so `truncated` is always `False` here; `test_ingest_
    scan_bounds.py::TestB5TheScanIsBoundedByPagesExamined` covers the
    `truncated=True` cases this class does not."""

    def test_returns_only_blank_page_numbers(self, tmp_path):
        path = tmp_path / "mixed.pdf"
        path.write_bytes(make_pdf_bytes(["a", None, "b", None]))

        assert readers.pdf_textless_pages(path) == ([2, 4], False)

    def test_scans_the_whole_document(self, tmp_path):
        """The exact fixture the retired `pdf_has_extractable_text`'s own
        `test_probe_pages_is_honored` used -- the regression this item
        exists to close: text sitting on page 6 must not be missed just
        because five textless pages come first."""
        path = tmp_path / "text-on-page-six.pdf"
        page_texts = [None, None, None, None, None, "text on six"]
        path.write_bytes(make_pdf_bytes(page_texts))

        assert readers.pdf_textless_pages(path) == ([1, 2, 3, 4, 5], False)

    def test_limit_stops_at_the_limit_th_textless_page(self, tmp_path):
        """`limit=1` over three fully textless pages must stop the SCAN
        after the first one is collected -- not scan all three and merely
        truncate the returned list (the load-bearing half: proven by
        counting `extract_text()` calls, not just checking the result)."""
        path = tmp_path / "three-textless.pdf"
        path.write_bytes(make_pdf_bytes([None, None, None]))

        original_extract_text = PageObject.extract_text
        with patch.object(PageObject, "extract_text", autospec=True, side_effect=original_extract_text) as mock_extract:
            found, truncated = readers.pdf_textless_pages(path, limit=1)

        assert found == [1]
        assert truncated is False
        assert mock_extract.call_count == 1

    def test_limit_still_scans_past_leading_text_pages_to_find_a_textless_one(self, tmp_path):
        """Orchestrator ruling R1: the original all-text `["a","b","c"]`
        fixture for this test was a tautology -- `limit` bounding pages
        COLLECTED vs. pages SCANNED makes no observable difference when
        there is nothing to collect at all, so it proved nothing about
        review N4's semantics. `["a","b",None]` does: the ONE textless
        page sits last, so a correct `limit=1` scan must still walk past
        BOTH text pages (three `extract_text()` calls) before it can
        collect page 3 and stop -- pinning that `limit` counts collected
        textless pages, never merely "the first `limit` pages scanned"."""
        path = tmp_path / "text-then-one-textless.pdf"
        path.write_bytes(make_pdf_bytes(["a", "b", None]))

        original_extract_text = PageObject.extract_text
        with patch.object(PageObject, "extract_text", autospec=True, side_effect=original_extract_text) as mock_extract:
            found, truncated = readers.pdf_textless_pages(path, limit=1)

        assert found == [3]
        assert truncated is False
        assert mock_extract.call_count == 3

    def test_limit_zero_returns_empty_list_without_opening_the_file(self, tmp_path):
        """W1 review MINOR 6: before this fix, `limit=0` still collected the
        FIRST textless page it found (nothing distinguished "0 collected"
        from "collect 0, then stop" in the loop's own `>= limit` check),
        silently behaving like `limit=1`. "Collect at most zero" has
        exactly one honest answer, `[]`, returned before a single page --
        or even the file itself -- is ever opened. Proven by pointing at a
        path with nothing on disk at all: a real scan attempt would raise."""
        path = tmp_path / "does-not-exist.pdf"

        assert readers.pdf_textless_pages(path, limit=0) == ([], False)

    def test_negative_limit_also_returns_empty_list(self, tmp_path):
        path = tmp_path / "does-not-exist.pdf"

        assert readers.pdf_textless_pages(path, limit=-1) == ([], False)
