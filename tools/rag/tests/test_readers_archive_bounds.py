"""B-6 (security audit round 3, H34): an archive-backed document
(`.docx`/`.xlsx`, both ZIP containers) is refused above a stated
uncompressed-size ceiling, before either parser ever expands it --
`readers.assert_archive_is_sane`, wired into `_read_docx` and the
`.xlsx` branch of `read_tabular_dataframe`. A `.csv` gets its own,
row-count-based bound (`readers.MAX_CSV_ROWS`) in the same function,
since it is not an archive and has no declared-vs-actual size to
compare.

**No bomb is ever decompressed here** (this host is recorded as
OOM-sensitive). Every fixture below is genuinely small on disk (at most
a couple of MB, written once, never expanded): the ZIP fixtures LIE
about their own uncompressed size by patching the central directory's
`file_size` field directly (`_patch_central_directory_uncompressed_size`)
rather than by actually writing that many bytes -- a real ZIP reader
trusts that field at listing time (`zipfile.ZipFile.infolist()`), which
is exactly the trust `assert_archive_is_sane` is built to be suspicious
of. Every refusal test also proves the parser was never reached, by
monkeypatching it to raise if called -- a guard that "parsed, then
noticed the problem" would already be too late.
"""
from __future__ import annotations

import struct
import zipfile
from pathlib import Path

import pytest

from tools.rag import readers

# A signature marking the start of one ZIP central-directory file header
# record (APPNOTE.TXT §4.3.12). Used only to locate the record this
# module patches -- never to build or read real archive contents.
_CENTRAL_DIR_SIGNATURE = b"PK\x01\x02"


def _patch_central_directory_uncompressed_size(path: Path, member: str, declared_size: int) -> None:
    """Rewrite the central directory record for `member` inside the ZIP
    at `path` so it DECLARES `declared_size` uncompressed bytes, without
    touching the member's actual (small) stored data at all -- the ZIP
    format keeps the uncompressed-size field in the central directory
    independent of the real bytes, which is exactly the gap
    `assert_archive_is_sane` exists to distrust. Never decompresses,
    never writes `declared_size` bytes anywhere; this function's own
    cost is one file read, one small in-place patch, one file write."""
    data = bytearray(path.read_bytes())
    offset = 0
    while True:
        idx = data.find(_CENTRAL_DIR_SIGNATURE, offset)
        if idx == -1:
            raise AssertionError(f"no central directory record found for {member!r}")
        name_len = struct.unpack_from("<H", data, idx + 28)[0]
        extra_len = struct.unpack_from("<H", data, idx + 30)[0]
        comment_len = struct.unpack_from("<H", data, idx + 32)[0]
        name = bytes(data[idx + 46 : idx + 46 + name_len]).decode("utf-8")
        if name == member:
            struct.pack_into("<I", data, idx + 24, declared_size)  # uncompressed size field
            path.write_bytes(bytes(data))
            return
        offset = idx + 46 + name_len + extra_len + comment_len


def _archive_with_lying_member(
    tmp_path: Path, filename: str, member: str, *, compressed_kb: int, declared_bytes: int
) -> Path:
    """Build a tiny, genuinely-valid ZIP at `tmp_path / filename` with one
    member (`member`, `compressed_kb` KiB of zero bytes, stored
    uncompressed so its real size on disk is exactly `compressed_kb *
    1024`), then patch that member's DECLARED uncompressed size to
    `declared_bytes` -- which may be arbitrarily larger without a single
    extra byte ever being written."""
    path = tmp_path / filename
    with zipfile.ZipFile(path, "w", zipfile.ZIP_STORED) as zf:
        zf.writestr(member, b"\x00" * (compressed_kb * 1024))
    _patch_central_directory_uncompressed_size(path, member, declared_bytes)
    return path


def ratio_bomb_docx(tmp_path: Path, *, compressed_kb: int, ratio: int) -> Path:
    """A `.docx`-suffixed ZIP whose `word/document.xml` member is really
    `compressed_kb` KiB on disk but declares `ratio` times that,
    uncompressed."""
    compressed_bytes = compressed_kb * 1024
    return _archive_with_lying_member(
        tmp_path, "bomb.docx", "word/document.xml",
        compressed_kb=compressed_kb, declared_bytes=compressed_bytes * ratio,
    )


def ratio_bomb_xlsx(tmp_path: Path, *, compressed_kb: int, ratio: int) -> Path:
    """The `.xlsx` sibling of `ratio_bomb_docx`, lying in
    `xl/worksheets/sheet1.xml` instead."""
    compressed_bytes = compressed_kb * 1024
    return _archive_with_lying_member(
        tmp_path, "bomb.xlsx", "xl/worksheets/sheet1.xml",
        compressed_kb=compressed_kb, declared_bytes=compressed_bytes * ratio,
    )


class TestB6ArchiveBackedDocumentsAreBounded:

    def test_a_high_ratio_document_is_refused_before_parsing(self, tmp_path, monkeypatch):
        """The fixture is ~40 KB compressed and declares ~40 MB
        uncompressed: over the ratio, under any allocation that matters
        if the guard regresses."""
        import docx

        path = ratio_bomb_docx(tmp_path, compressed_kb=40, ratio=1000)

        def _must_not_be_called(*args, **kwargs):
            raise AssertionError("docx.Document was called -- the guard let a bomb through")

        monkeypatch.setattr(docx, "Document", _must_not_be_called)

        with pytest.raises(readers.ArchiveExpansionExceededError) as excinfo:
            readers.read_prose_documents(path)
        assert "compresses too far" in str(excinfo.value)

    def test_the_new_type_is_a_value_error_so_existing_handlers_still_see_it(self):
        assert issubclass(readers.ArchiveExpansionExceededError, ValueError)

    def test_an_absolute_ceiling_refuses_a_low_ratio_giant(self, tmp_path, monkeypatch):
        """A ratio comfortably under `MAX_UNCOMPRESSED_RATIO` still trips
        the absolute-bytes ceiling once the DECLARED total is large
        enough -- the case the ratio check alone would miss."""
        import docx

        assert readers.MAX_UNCOMPRESSED_BYTES < 1024 * 1024 * 1024  # sanity: this test's arithmetic below assumes it
        compressed_kb = 2048  # 2 MB real, on-disk data -- still "a few megabytes"
        ratio = readers.MAX_UNCOMPRESSED_RATIO - 50  # comfortably under the ratio ceiling
        declared = compressed_kb * 1024 * ratio
        assert declared > readers.MAX_UNCOMPRESSED_BYTES  # the fixture must actually exercise this branch

        path = _archive_with_lying_member(
            tmp_path, "giant.docx", "word/document.xml",
            compressed_kb=compressed_kb, declared_bytes=declared,
        )

        monkeypatch.setattr(
            docx, "Document",
            lambda *a, **k: (_ for _ in ()).throw(AssertionError("docx.Document was called")),
        )

        with pytest.raises(readers.ArchiveExpansionExceededError) as excinfo:
            readers.read_prose_documents(path)
        assert "uncompressed bytes total" in str(excinfo.value)

    def test_an_ordinary_document_parses_unchanged(self, tmp_path):
        """A real, small .docx/.xlsx/.csv -- nowhere near either
        ceiling -- still parses exactly as before this guard existed."""
        import docx
        import openpyxl

        docx_path = tmp_path / "ordinary.docx"
        d = docx.Document()
        d.add_paragraph("hello from an ordinary document")
        d.save(docx_path)
        docs = readers.read_prose_documents(docx_path)
        assert len(docs) == 1
        assert "hello from an ordinary document" in docs[0].text

        xlsx_path = tmp_path / "ordinary.xlsx"
        wb = openpyxl.Workbook()
        wb.active["A1"] = "hello"
        wb.save(xlsx_path)
        df = readers.read_tabular_dataframe(xlsx_path)
        assert list(df.columns) == ["hello"]

        csv_path = tmp_path / "ordinary.csv"
        csv_path.write_text("a,b\n1,2\n3,4\n", encoding="utf-8")
        df = readers.read_tabular_dataframe(csv_path)
        assert len(df) == 2

    def test_a_spreadsheet_is_checked_the_same_way(self, tmp_path, monkeypatch):
        """Same ratio-ceiling refusal, for `.xlsx` -- `read_tabular_dataframe`
        never reaches `pandas.read_excel`."""
        import pandas as pd

        path = ratio_bomb_xlsx(tmp_path, compressed_kb=40, ratio=1000)

        def _must_not_be_called(*args, **kwargs):
            raise AssertionError("pandas.read_excel was called -- the guard let a bomb through")

        monkeypatch.setattr(pd, "read_excel", _must_not_be_called)

        with pytest.raises(readers.ArchiveExpansionExceededError) as excinfo:
            readers.read_tabular_dataframe(path)
        assert "compresses too far" in str(excinfo.value)

    def test_a_csv_is_bounded_by_rows(self, tmp_path, monkeypatch):
        """Same shape for a highly compressible CSV, which is not an
        archive and therefore needs its own bound."""
        import pandas as pd

        # Lower the ceiling rather than writing a million-row fixture --
        # the row-count arithmetic this test proves is independent of
        # where the ceiling actually sits.
        monkeypatch.setattr(readers, "MAX_CSV_ROWS", 5)

        def _must_not_be_called(*args, **kwargs):
            raise AssertionError("pandas.read_csv was called -- the guard let a bomb through")

        monkeypatch.setattr(pd, "read_csv", _must_not_be_called)

        path = tmp_path / "big.csv"
        lines = ["a"] + [str(i) for i in range(6)]  # header + 6 data rows, over the lowered cap of 5
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")

        with pytest.raises(readers.ArchiveExpansionExceededError) as excinfo:
            readers.read_tabular_dataframe(path)
        assert "row" in str(excinfo.value)

    def test_the_guard_never_expands_the_archive_to_measure_it(self, tmp_path, monkeypatch):
        """It reads the central directory's declared sizes. A guard that
        decompressed to measure would BE the bomb."""
        path = ratio_bomb_docx(tmp_path, compressed_kb=40, ratio=1000)

        def _must_not_be_called(*args, **kwargs):
            raise AssertionError("zipfile.ZipFile.read was called -- the guard decompressed to measure")

        monkeypatch.setattr(zipfile.ZipFile, "read", _must_not_be_called)

        with pytest.raises(readers.ArchiveExpansionExceededError):
            readers.assert_archive_is_sane(path)


def _archive_with_lying_members(tmp_path: Path, filename: str, members) -> Path:
    """`_archive_with_lying_member`'s multi-member sibling: `members` is a
    sequence of `(member_name, compressed_kb, declared_bytes)` triples,
    each written as `compressed_kb` KiB of real zero bytes and then
    patched to DECLARE `declared_bytes` uncompressed. Same guarantee --
    nothing is ever expanded, and the archive on disk stays the sum of
    the `compressed_kb` values."""
    path = tmp_path / filename
    with zipfile.ZipFile(path, "w", zipfile.ZIP_STORED) as zf:
        for member, compressed_kb, _declared in members:
            zf.writestr(member, b"\x00" * (compressed_kb * 1024))
    for member, _compressed_kb, declared in members:
        _patch_central_directory_uncompressed_size(path, member, declared)
    return path


class TestB6TheSpreadsheetSumCountsOnlyWhatOpenpyxlOpens:
    """H34 round-3 follow-up, a NARROW false positive: the absolute
    ceiling summed EVERY member of an `.xlsx`, including the parts
    openpyxl's `read_only` mode never opens at all.

    `pd.read_excel(..., engine_kwargs={"read_only": True})` streams the
    worksheet XML and reads the workbook, shared strings, styles and
    relationship parts. It does not touch `xl/media/*` (embedded
    pictures) or `xl/drawings/*` (their anchors and shapes) -- those are
    parts the drawing model would need, and read-only mode never builds
    one. A legitimate workbook carrying a few hundred megabytes of
    photographs was therefore refused for an expansion that never
    happens.

    `.docx` keeps whole-archive accounting: python-docx builds the full
    OPC package, embedded media included, so every member really is
    loaded there."""

    # A member 2 MB on disk declaring a ratio comfortably UNDER
    # `MAX_UNCOMPRESSED_RATIO`, so it is the absolute-bytes ceiling these
    # tests exercise and never the per-member ratio one (the existing
    # `test_an_absolute_ceiling_refuses_a_low_ratio_giant` above picks its
    # own numbers the same way, for the same reason).
    _GIANT_KB = 2048
    _GIANT_RATIO = readers.MAX_UNCOMPRESSED_RATIO - 50
    _GIANT_DECLARED = _GIANT_KB * 1024 * _GIANT_RATIO

    def test_the_fixture_really_exercises_the_absolute_ceiling(self):
        """A sanity pin, not a behaviour: the arithmetic every test below
        depends on."""
        assert self._GIANT_DECLARED > readers.MAX_UNCOMPRESSED_BYTES
        assert self._GIANT_RATIO < readers.MAX_UNCOMPRESSED_RATIO

    def test_a_workbook_with_a_large_embedded_image_is_not_refused(self, tmp_path, monkeypatch):
        """The false positive itself: the picture alone is over the
        absolute ceiling, and the sheet openpyxl actually streams is
        tiny. The guard must let this through to the parser."""
        import pandas as pd

        path = _archive_with_lying_members(tmp_path, "photos.xlsx", [
            ("xl/worksheets/sheet1.xml", 1, 4096),
            ("xl/media/image1.png", self._GIANT_KB, self._GIANT_DECLARED),
        ])

        sentinel = object()
        monkeypatch.setattr(pd, "read_excel", lambda *a, **k: sentinel)
        assert readers.read_tabular_dataframe(path) is sentinel

    def test_a_drawing_part_is_excluded_the_same_way(self, tmp_path, monkeypatch):
        """`xl/drawings/*` is the picture's anchor, and read-only mode
        builds no drawing model to need it."""
        import pandas as pd

        path = _archive_with_lying_members(tmp_path, "shapes.xlsx", [
            ("xl/worksheets/sheet1.xml", 1, 4096),
            ("xl/drawings/drawing1.xml", self._GIANT_KB, self._GIANT_DECLARED),
        ])

        sentinel = object()
        monkeypatch.setattr(pd, "read_excel", lambda *a, **k: sentinel)
        assert readers.read_tabular_dataframe(path) is sentinel

    def test_the_same_shape_in_a_document_is_still_refused(self, tmp_path, monkeypatch):
        """The other half of the ruling: python-docx loads every part, so
        `.docx` still sums the whole archive and this refuses."""
        import docx

        path = _archive_with_lying_members(tmp_path, "photos.docx", [
            ("word/document.xml", 1, 4096),
            ("word/media/image1.png", self._GIANT_KB, self._GIANT_DECLARED),
        ])

        monkeypatch.setattr(
            docx, "Document",
            lambda *a, **k: (_ for _ in ()).throw(AssertionError("docx.Document was called")),
        )

        with pytest.raises(readers.ArchiveExpansionExceededError) as excinfo:
            readers.read_prose_documents(path)
        assert "uncompressed bytes total" in str(excinfo.value)

    def test_an_oversize_worksheet_is_still_refused_in_a_workbook_with_media(
        self, tmp_path, monkeypatch
    ):
        """The narrowing must not become a bypass: a bomb that hides in
        the parts openpyxl DOES open still trips the ceiling, however
        much unread media sits beside it."""
        import pandas as pd

        path = _archive_with_lying_members(tmp_path, "bomb.xlsx", [
            ("xl/worksheets/sheet1.xml", self._GIANT_KB, self._GIANT_DECLARED),
            ("xl/media/image1.png", 1, 4096),
        ])

        def _must_not_be_called(*args, **kwargs):
            raise AssertionError("pandas.read_excel was called -- the guard let a bomb through")

        monkeypatch.setattr(pd, "read_excel", _must_not_be_called)

        with pytest.raises(readers.ArchiveExpansionExceededError) as excinfo:
            readers.read_tabular_dataframe(path)
        assert "uncompressed bytes total" in str(excinfo.value)

    def test_the_ratio_ceiling_still_covers_every_member(self, tmp_path, monkeypatch):
        """Only the SUM narrowed. A member that compresses 1000:1 is a
        crafted member wherever it sits -- a real picture or drawing part
        is already compressed and never approaches the ratio -- so the
        per-member check stays whole-archive."""
        import pandas as pd

        compressed_bytes = 40 * 1024
        path = _archive_with_lying_members(tmp_path, "ratio.xlsx", [
            ("xl/worksheets/sheet1.xml", 1, 4096),
            ("xl/media/image1.png", 40, compressed_bytes * 1000),
        ])

        def _must_not_be_called(*args, **kwargs):
            raise AssertionError("pandas.read_excel was called -- the guard let a bomb through")

        monkeypatch.setattr(pd, "read_excel", _must_not_be_called)

        with pytest.raises(readers.ArchiveExpansionExceededError) as excinfo:
            readers.read_tabular_dataframe(path)
        assert "compresses too far" in str(excinfo.value)


class TestB6AnEmptyCsvCountsZeroRows:
    """`_count_csv_rows` subtracts the header line unconditionally, so a
    genuinely empty file counted -1 rows. Nothing downstream refused on
    it (the bound is an upper one), but a negative row count is not a
    fact about any file, and it is the value `read_tabular_dataframe`
    compares against the ceiling."""

    def test_an_empty_file_counts_zero_not_minus_one(self, tmp_path):
        path = tmp_path / "empty.csv"
        path.write_text("", encoding="utf-8")
        assert readers._count_csv_rows(path) == 0

    def test_a_header_only_file_still_counts_zero(self, tmp_path):
        path = tmp_path / "header-only.csv"
        path.write_text("a,b\n", encoding="utf-8")
        assert readers._count_csv_rows(path) == 0
